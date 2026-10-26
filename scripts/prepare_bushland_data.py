"""Extract observed USDA Bushland records without modifying source workbooks.

Run with the bundled Python/openpyxl/pandas runtime. Missing Excel markers remain
missing; observed dry grain is distinct from standard-moisture yield. Source
sheet names and Excel row numbers accompany every biological measurement.
"""
from __future__ import annotations
import json
import math
import re
import sys
from datetime import datetime,timedelta
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from verify_bushland_sources import verify

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'data/external/bushland'
OUT=ROOT/'data/bushland_processed'
YEARS=[1989,1990,1994,2013,2016,2018]


def number(value):
    if isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value):return float(value)
    return np.nan


def clean(text):return re.sub(r'\s+',' ',str(text or '').lower()).strip()


def rows(sheet):
    iterator=sheet.iter_rows(max_col=min(sheet.max_column,40),values_only=True)
    headers=next(iterator)
    headers=[clean(h) for h in headers]
    for index,row in enumerate(iterator,2):
        if any(x is not None for x in row):yield index,dict(zip(headers,row))


def get(record,patterns):
    for pattern in patterns:
        for key,value in record.items():
            if re.fullmatch(pattern,key):return value
    return None


def main():
    verify()
    OUT.mkdir(parents=True,exist_ok=True)
    calendars=[];notes=[]
    for path in sorted((RAW/'24855873').glob('*.xlsx')):
        year=int(path.name[:4]);side='E' if '_East_' in path.name else 'W'
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        data=next(s for s in w if not s.title.lower().startswith('dic'))
        events=[]
        for i,row in enumerate(data.iter_rows(max_col=8,values_only=True),1):
            texts=[str(c) for c in row if isinstance(c,str)]
            event=' '.join(texts)
            date=next((c for c in row if isinstance(c,datetime)),None)
            if date and event:events.append((date,i,event))
        plants=[x for x in events if re.search(r'\b(?:planted|started planting)\b',x[2],re.I)
                and not re.search(r'irrigation|hand.planted skips|replanted',x[2],re.I)]
        if not plants:raise ValueError(f'No planting event: {path.name}')
        for field in (['NE','SE'] if side=='E' else ['NW','SW']):
            specific=[x for x in plants if re.search(r'\b'+field+r'\b',x[2],re.I)]
            planting=min(specific or plants,key=lambda r:r[0]);plant_doy=planting[0].timetuple().tm_yday
            calendars.append({'year':year,'field':field,'planting_date':planting[0].date().isoformat(),
                              'planting_doy':plant_doy,'source_file':path.name,'source_sheet':data.title,
                              'source_row':planting[1],'source_event':planting[2]})
        w.close()
    calendar=pd.DataFrame(calendars)
    if len(calendar)!=20:raise ValueError(f'Expected 20 field-seasons, got {len(calendar)}')
    calendar.to_csv(OUT/'calendars.csv',index=False)

    weather=[]
    reuse_weather='--reuse-weather' in sys.argv and (OUT/'daily_weather.csv').exists()
    for year in ([] if reuse_weather else YEARS):
        path=next((RAW/'25316164').glob(f'{year}_*.xlsx'))
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        s=w.worksheets[2]
        daily={}
        for _,record in rows(s):
            doy=number(get(record,['day of year']))
            if not np.isfinite(doy):continue
            key=int(doy);values=daily.setdefault(key,[])
            values.append({'temp':number(get(record,['air temperature in °c','air temperature in c'])),
                           'solar':number(get(record,['solar irradiance in w/m2',r'solar irradiance in w/m\^2'])),
                           'rh':number(get(record,['rh in %'])),
                           'wind':number(get(record,['wind speed in m/s'])),
                           'pressure':number(get(record,['air pressure in kpa'])),
                           'rain':number(get(record,['average precipitation in mm']))})
        for doy,records in sorted(daily.items()):
            frame=pd.DataFrame(records)
            counts=frame.count()
            rec={'year':year,'doy':doy,'intervals':len(frame),
                 'temp_valid':int(counts['temp']),'solar_valid':int(counts['solar']),
                 'tmean_c':frame.temp.mean(),'tmin_c':frame.temp.min(),'tmax_c':frame.temp.max(),
                 'solar_mj_m2':frame.solar.clip(lower=0).sum(min_count=96)*900/1e6,
                 'rh_mean_pct':frame.rh.mean(),'wind_m_s':frame.wind.mean(),'pressure_kpa':frame.pressure.mean(),
                 'rain_gauge_mm':frame.rain.sum(min_count=96),
                 'quality_complete':bool(len(frame)==96 and (counts[['temp','solar','rh','wind']]==96).all()),
                 'pressure_measured':bool(counts['pressure']==96),
                 'source_file':path.name,'source_sheet':s.title}
            weather.append(rec)
        w.close()
        print(f'Weather {year}: {len(daily)} daily records',flush=True)
    if weather:pd.DataFrame(weather).to_csv(OUT/'daily_weather.csv',index=False)

    water=[]
    for path in sorted((RAW/'24855957').glob('*.xlsx')):
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        s=next(s for s in w if 'daily wbal' in s.title.lower() and not s.title.lower().startswith('dic'))
        for i,r in rows(s):
            year=number(get(r,['year']));doy=number(get(r,['day of year']))
            if not np.isfinite(year+doy):continue
            for field in ['NE','SE','NW','SW']:
                prefix=field.lower()
                rec={'year':int(year),'doy':int(doy),'field':field,'source_file':path.name,
                     'source_sheet':s.title,'source_row':i}
                for label,pattern in [('et_mm','et'),('dew_mm','dew'),('rain_mm','precip'),
                                      ('irrigation_mm',r'irrig\.'),('counterweight_mm',r'cw adjust\.'),('drainage_signed_mm','drainage')]:
                    rec[label]=number(get(r,[f'{prefix} {pattern} in mm']))
                if np.isfinite(rec['et_mm']):water.append(rec)
        # A separate official event log retains irrigation on dates where the
        # daily lysimeter column is missing. Do not rewrite missing measurements.
        irrigation_sheet=next(s for s in w if not s.title.lower().startswith('dic')
                              and 'irrig' in s.title.lower() and 'daily' in s.title.lower())
        events=[]
        for i,r in rows(irrigation_sheet):
            doy=number(get(r,['doy','day of year']));year=number(get(r,['year']))
            if not np.isfinite(doy+year):continue
            for field in ['NE','SE','NW','SW']:
                value=number(get(r,[field.lower()+' irrigation in mm']))
                if np.isfinite(value):events.append({'year':int(year),'doy':int(doy),'field':field,
                    'irrigation_mm':value,'source_file':path.name,'source_sheet':irrigation_sheet.title,'source_row':i})
        pd.DataFrame(events).to_csv(OUT/f'irrigation_events_{int(path.name[:4])}.csv',index=False)
        w.close()
    pd.DataFrame(water).to_csv(OUT/'daily_water_observed.csv',index=False)

    growth=[];yields=[]
    for path in sorted((RAW/'24855909').glob('*.xlsx')):
        year=int(path.name[:4]);side='E' if '_East_' in path.name else 'W'
        w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        for s in w:
            name=s.title.lower()
            if name.startswith(('dic','dict')) or 'intro' in name:continue
            if ('comb' in name and ('harv' in name or 'yield' in name)):
                for i,r in rows(s):
                    span=number(get(r,['span number',r'span no\.','span','plot number']))
                    sample_id=str(get(r,['sample id']) or '').upper().strip()
                    if not np.isfinite(span) and re.fullmatch(r'(NE|SE|NW|SW)\d+',sample_id):
                        span=int(re.search(r'\d+',sample_id).group())
                    if not np.isfinite(span) or not 1<=span<=10:continue
                    field=('N' if span<=5 else 'S')+side
                    source_field=str(get(r,['field','location']) or '').strip().upper()
                    if source_field in ['NE','SE','NW','SW']:field=source_field
                    elif source_field:continue  # e.g. separate zinc treatment strips
                    elif sample_id[:2] in ['NE','SE','NW','SW']:field=sample_id[:2]
                    dry=number(get(r,['dry yield in kg/ha','yield of dry grain in kg/ha','dry grain yield in kg/ha']))
                    dry_kg_m2=number(get(r,[r'dry yield in kg/m\^2']))
                    if np.isfinite(dry_kg_m2):dry=dry_kg_m2*10000.
                    standard_bu=number(get(r,['yield in bu/acre at standard moisture','yield at standard moisture in bu/acre',
                                                'grain yield at standard moisture in bu/acre']))
                    area=number(get(r,[r'sample area in m\^2',r'number of m\^2']))
                    if not np.isfinite(area):
                        acres=number(get(r,['harvested area in acres','number of acres in sample','number of acres','plot area in acres']))
                        area=acres*4046.8564224
                    if np.isfinite(dry):
                        yields.append({'year':year,'field':field,'span':int(span),'dry_grain_kg_ha':dry,
                            'standard_moisture_kg_ha':standard_bu*62.77,'harvested_area_m2':area,
                            'irrigation_treatment_pct':number(get(r,['irrigation treatment in percent'])),
                            'source_file':path.name,'source_sheet':s.title,'source_row':i,
                            'date_year_inconsistent':isinstance(get(r,['date']),datetime) and get(r,['date']).year!=year})
            # Explicit LAI/biomass tabs for 2013 onward, combined Growth tabs before then.
            if 'lai' in name and 'biomass' in name or (year<2000 and name.endswith('maize growth')):
                for i,r in rows(s):
                    date=get(r,['date']);doy=number(get(r,['doy','day of year']))
                    if not np.isfinite(doy) and isinstance(date,datetime):doy=date.timetuple().tm_yday
                    if not np.isfinite(doy):continue
                    span=number(get(r,['span number',r'span no\.']))
                    field_text=str(get(r,['field','location','plot id']) or '').upper().strip()
                    if field_text[:2] in ['NE','SE','NW','SW']:field=field_text[:2]
                    elif np.isfinite(span):field=('N' if span<=5 else 'S')+side
                    else:
                        plot=number(get(r,['plot number','sample id']))
                        if not np.isfinite(plot):continue
                        field=('N' if plot<=3 else 'S')+side
                    lai=number(get(r,['lai','leaf area index']))
                    biomass=number(get(r,[r'.*above.ground dry matter.*g/m\^2','.*above.ground dry matter.*g/m2',
                                           r'total biomass in g/m\^2',r'total dry matter in g/m\^2']))
                    biomass_kg_ha=number(get(r,['.*above.ground dry matter.*kg/ha']))
                    if np.isfinite(biomass_kg_ha):biomass=biomass_kg_ha/10.
                    if not np.isfinite(biomass):
                        # Do not manufacture biomass from incomplete organ masses.
                        mass=number(get(r,['total dry mass in g','total above.ground dry mass in g','total above.ground dry biomass in g']))
                        area=number(get(r,[r'plot size in m\^2']))
                        if np.isfinite(mass) and np.isfinite(area) and area>0:biomass=mass/area
                    if np.isfinite(lai) or np.isfinite(biomass):
                        growth.append({'year':year,'field':field,'doy':int(doy),'lai_observed':lai,
                                       'biomass_kg_ha':biomass*10,'source_file':path.name,
                                       'source_sheet':s.title,'source_row':i})
        w.close()
    pd.DataFrame(yields).to_csv(OUT/'grain_yield_observed.csv',index=False)
    pd.DataFrame(growth).to_csv(OUT/'growth_observed.csv',index=False)

    profiles=[]
    for path in sorted((RAW/'24855858').glob('*Maize_Soil-water.xlsx')):
        year=int(path.name[:4]);w=openpyxl.load_workbook(path,read_only=True,data_only=True)
        s=next(s for s in w if 'soil vwc' in s.title.lower() and not s.title.lower().startswith('dic'))
        for i,r in rows(s):
            doy=number(get(r,['doy','day of year']));depth=number(get(r,['depth in cm','depth']))
            vwc=number(get(r,['vwc','water content in m3/m3']))
            location=str(get(r,['location']) or '')
            if not np.isfinite(doy+depth+vwc) or not re.match(r'^(NE|SE|NW|SW)',location):continue
            if not 0<=vwc<=.65:continue
            profiles.append({'year':year,'field':location[:2],'doy':int(doy),'depth_cm':depth,
                'theta_m3_m3':vwc,'location':location,'is_lysimeter':'lys' in location.lower(),
                'source_file':path.name,'source_sheet':s.title,'source_row':i})
        w.close()
    pd.DataFrame(profiles).to_csv(OUT/'soil_water_observed.csv',index=False)
    summary={'years':YEARS,'field_seasons':len(calendar),'yield_observations':len(yields),
             'growth_observations':len(growth),'water_days':len(water),'soil_water_observations':len(profiles),
             'sources':'USDA Ag Data Commons; original raw workbooks preserved with acquisition checksums',
             'missing':'#N/A, blanks, nonnumeric markers retained as missing; no zero substitution',
             'weather':'96 original 15-minute records per source day; radiation W/m2 integrated to MJ/m2/day; no imputation',
             'grain':'dry grain kg/ha is not equal to standard-moisture grain kg/ha',
             'field_mapping':'source field when present; otherwise spans 1-5 north, 6-10 south, east/west workbook side',
             'causal_limit':'75% and 100% sprinkler treatment groups occupy distinct NW/SW fields; field and irrigation are confounded'}
    (OUT/'preparation_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary,indent=2))


if __name__=='__main__':main()
