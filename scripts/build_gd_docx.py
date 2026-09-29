from pathlib import Path
import hashlib
import json
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def transcribir_backend_a_word():
    """Transcribe íntegramente los archivos de texto del backend a GD.docx."""
    paths = [str(p.relative_to(ROOT)) for p in (ROOT / 'backend').rglob('*')
             if p.is_file() and not any(part in {'__pycache__', '.pytest_cache', '.venv', 'venv'}
                                        for part in p.relative_to(ROOT).parts)]
    sources = []
    for name in sorted(paths):
        path = ROOT / name
        if path.suffix.lower() not in {'.py', '.json', '.txt', '.sh', '.md', '.example', '.yaml', '.yml', '.toml', '.sql'}:
            continue
        content = path.read_bytes().decode('utf-8-sig')
        sources.append((name.replace('\\', '/'), content, path))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    code = doc.styles.add_style('Codigo fuente', 1)
    code.font.name = 'Consolas'
    code.font.size = Pt(8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.widow_control = False
    doc.add_paragraph('Código del backend del gemelo digital CeresPINN', 'Title')
    doc.add_paragraph('Transcripción completa de los archivos de texto del backend, incluidos API, inferencia, entrenamiento, proveedores de datos, pruebas, dependencias y configuración. Cada sección identifica la ruta del archivo y conserva su contenido e indentación.')
    doc.add_paragraph(f'Contenido: {len(sources)} archivos de texto. El archivo backend/models/cerespinn_pinn.pt contiene pesos binarios del modelo y no corresponde a código fuente transcribible.')
    doc.add_heading('Índice de archivos', 1)
    for name, _, _ in sources:
        p = doc.add_paragraph(name)
        p.paragraph_format.space_after = Pt(1)
        p.runs[0].font.size = Pt(9)
    expected = []
    for name, content, path in sources:
        doc.add_page_break()
        doc.add_heading('Archivo del backend', 1)
        doc.add_paragraph(name)
        normalized = content.replace('\r\n', '\n').replace('\r', '\n')
        for line in normalized.split('\n'):
            doc.add_paragraph(line, 'Codigo fuente')
        expected.append({'ruta': name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                         'lineas': len(normalized.split('\n'))})
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Página ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    target = ROOT / 'GD.docx'
    doc.save(target)
    # Comprueba el contenido completo después de volver a abrir el documento.
    loaded = Document(target)
    paras = iter(loaded.paragraphs)
    found = 0
    for p in paras:
        if p.text == 'Archivo del backend' and p.style.name == 'Heading 1':
            route = next(paras).text
            name, content, _ = sources[found]
            assert route == name
            lines = content.replace('\r\n', '\n').replace('\r', '\n').split('\n')
            actual = [next(paras).text for _ in lines]
            assert actual == lines, name
            found += 1
    assert found == len(sources)
    qa = ROOT / '.gd_qa'
    qa.mkdir(exist_ok=True)
    (qa / 'manifest.json').write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Verificado: {found} archivos, {sum(x["lineas"] for x in expected)} líneas. {target}')

if __name__ == '__main__':
    transcribir_backend_a_word()
