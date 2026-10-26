"""Build the corrected English manuscript from verified local evidence.

Uses the supplied Word document as the typography/page-layout reference. The
original is never overwritten. All numerical tables are read from audit outputs.
"""
from pathlib import Path
import csv
import hashlib
import json
import re
import shutil
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'manuscript_source.docx'
if not SOURCE.exists():
    SOURCE=Path(r'C:\Users\USERJSSV\Downloads\CeresPINN_LucanoNieves_RojasVillgas.docx')
ART=ROOT/'docs/q1_artifacts'
REV=ART/'revision_20261026'
DEST=ROOT/'deliverables/revision_20261026'

# Numbering is assigned explicitly and checked bidirectionally at the end.
REFERENCES=[
('Xiao, K., et al. (2025). Drought and extreme heat reduce wheat and maize production in the United States by lowering both crop yields and harvestable fraction. Earth\u2019s Future, 13, e2024EF005557.','10.1029/2024EF005557'),
('Jägermeyr, J., et al. (2021). Climate impacts on global agriculture emerge earlier in new generation of climate and crop models. Nature Food, 2, 873–885.','10.1038/s43016-021-00400-y'),
('Xu, Q., et al. (2024). Assessing climate change impacts on crop yields and exploring adaptation strategies in Northeast China. Earth\u2019s Future, 12, e2023EF004063.','10.1029/2023EF004063'),
('Shahhosseini, M., Hu, G., Huber, I., and Archontoulis, S. V. (2021). Coupling machine learning and crop modeling improves crop yield prediction in the US Corn Belt. Scientific Reports, 11, 1606.','10.1038/s41598-020-80820-1'),
('Cheng, M., et al. (2022). Combining multi-indicators with machine-learning algorithms for maize yield early prediction at the county-level in China. Agricultural and Forest Meteorology, 323, 109057.','10.1016/j.agrformet.2022.109057'),
('Escribà-Gelonch, M., et al. (2024). Digital twins in agriculture: Orchestration and applications. Journal of Agricultural and Food Chemistry, 72, 10737–10752.','10.1021/acs.jafc.4c01934'),
('Zhang, R., Zhu, H., Chang, Q., and Mao, Q. (2025). A comprehensive review of digital twins technology in agriculture. Agriculture, 15, 903.','10.3390/agriculture15090903'),
('Melesse, T. Y. (2025). Digital twin-based applications in crop monitoring. Heliyon, 11, e42137.','10.1016/j.heliyon.2025.e42137'),
('Raissi, M., Perdikaris, P., and Karniadakis, G. E. (2019). Physics-informed neural networks: A deep learning framework for solving forward and inverse problems involving nonlinear partial differential equations. Journal of Computational Physics, 378, 686–707.','10.1016/j.jcp.2018.10.045'),
('Shi, Y., et al. (2025). Deep learning meets process-based models: A hybrid approach to agricultural challenges. arXiv preprint, arXiv:2504.16141.','10.48550/arXiv.2504.16141'),
('Shi, Y., et al. (2026). AgriPINN: A process-informed neural network for interpretable and scalable crop biomass prediction under water stress. arXiv preprint, arXiv:2601.16045.','10.48550/arXiv.2601.16045'),
('Chavoshi, A., Dashtian, H., Bakhshian, S., Young, M. H., and Niyogi, D. (2025). PINN-SM: A physics-informed neural networks model for vadose zone soil moisture profile prediction. Journal of Geophysical Research: Machine Learning and Computation, 2, e2024JH000547.','10.1029/2024JH000547'),
('Xu, X., et al. (2025). Digital twin-based winter wheat growth simulation and optimization. Field Crops Research, 329, 109953.','10.1016/j.fcr.2025.109953'),
('Xu, C., Torres-Rojas, L., Vergopolan, N., and Chaney, N. W. (2023). The benefits of using state-of-the-art digital soil properties maps to improve the modeling of soil moisture in land surface models. Water Resources Research, 59, e2022WR032336.','10.1029/2022WR032336'),
('Purcell, W., Neubauer, T., and Mallinger, K. (2023). Digital twins in agriculture: Challenges and opportunities for environmental sustainability. Current Opinion in Environmental Sustainability, 61, 101252.','10.1016/j.cosust.2022.101252'),
('Eyring, V., et al. (2016). Overview of the Coupled Model Intercomparison Project Phase 6 (CMIP6) experimental design and organization. Geoscientific Model Development, 9, 1937–1958.','10.5194/gmd-9-1937-2016'),
('Thrasher, B., Wang, W., Michaelis, A., Melton, F., Lee, T., and Nemani, R. (2022). NASA global daily downscaled projections, CMIP6. Scientific Data, 9, 262.','10.1038/s41597-022-01393-4'),
('Maraun, D., et al. (2010). Precipitation downscaling under climate change: Recent developments to bridge the gap between dynamical models and the end user. Reviews of Geophysics, 48, RG3003.','10.1029/2009RG000314'),
('Mbengue, A., et al. (2026). High-resolution downscaled CMIP6 projections dataset of key climate variables for Senegal. Scientific Data, 13, 723.','10.1038/s41597-026-07059-9'),
('Galmarini, S., et al. (2024). Assessing the impact on crop modelling of multi- and uni-variate climate model bias adjustments. Agricultural Systems, 215, 103846.','10.1016/j.agsy.2023.103846'),
('O\u2019Neill, B. C., et al. (2016). The Scenario Model Intercomparison Project (ScenarioMIP) for CMIP6. Geoscientific Model Development, 9, 3461–3482.','10.5194/gmd-9-3461-2016'),
('Minoli, S., Jägermeyr, J., Asseng, S., Urfels, A., and Müller, C. (2022). Global crop yields can be lifted by timely adaptation of growing periods to climate change. Nature Communications, 13, 7079.','10.1038/s41467-022-34411-5'),
('Abramoff, R. Z., et al. (2023). Adaptation strategies strongly reduce the future impacts of climate change on simulated crop yields. Earth\u2019s Future, 11, e2022EF003190.','10.1029/2022EF003190'),
('Zhang, L., et al. (2025). Increased irrigation could mitigate future warming-induced maize yield losses in the Ogallala Aquifer. Communications Earth & Environment, 6, 483.','10.1038/s43247-025-02459-y'),
('Alimagham, S., et al. (2025). Integrating crop models and machine learning for projecting climate change impacts on crops in data-limited environments. Agricultural Systems, 228, 104367.','10.1016/j.agsy.2025.104367'),
('NOAA National Centers for Environmental Information. Climate at a Glance. County temperature and precipitation time series. Accessed October 26, 2026.','https://www.ncei.noaa.gov/access/monitoring/climate-at-a-glance/'),
('USDA National Agricultural Statistics Service. Quick Stats. Maize yield statistics. Accessed October 26, 2026.','https://quickstats.nass.usda.gov/'),
('Evett, S. R., Marek, G. W., Colaizzi, P. D., Copeland, K. S., Ruthardt, B. B., and Howell, T. A. (2025). The Bushland, Texas, maize evapotranspiration, growth, and yield dataset Collection. Scientific Data, 12, 209.','10.1038/s41597-025-04539-2'),
('Allen, R. G., Pereira, L. S., Raes, D., and Smith, M. (1998). Crop evapotranspiration: Guidelines for computing crop water requirements. FAO Irrigation and Drainage Paper 56.','https://www.fao.org/4/X0490E/X0490E00.htm'),
('USDA Agricultural Research Service. Bushland maize growth, soil-water, weather, agronomic calendar, and water-balance datasets. Ag Data Commons. Collection DOI; individual dataset DOIs and file checksums in the acquisition manifest.','10.15482/USDA.ADC/1526317'),
('USDA National Agricultural Statistics Service. Iowa agricultural statistics districts. Crop reporting district map and county listing. Accessed October 26, 2026.','https://www.nass.usda.gov/Statistics_by_State/Iowa/Publications/Crop_Report/IA_Ag_Districts_Map.pdf'),
]


def loadcsv(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def main():
    DEST.mkdir(parents=True,exist_ok=True)
    audit=json.loads((REV/'audit.json').read_text())
    stats=json.loads((ART/'q1_statistics.json').read_text())
    temporal=json.loads((ART/'temporal/temporal_validation_1990_2017_to_2018_2025.json').read_text())
    ext=json.loads((ART/'extended/extended_audit.json').read_text())
    field_metrics=json.loads((ART/'bushland/external_metrics.json').read_text())
    reacquired=json.loads((ART/'reacquired_20261026/reacquired_county_benchmark.json').read_text())
    districts=json.loads((ART/'reacquired_20261026/district_benchmark.json').read_text())
    baseline=loadcsv(ART/'extended/baseline_metrics_full_precision.csv')
    multi=loadcsv(ART/'extended/multiseed_neural_summary.csv')
    manage=loadcsv(ART/'extended/management_scenarios.csv')
    scores={r['model']:r for r in audit['national']['scores']}
    d=Document(SOURCE)
    body=d._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    # Retain the supplied A4 page geometry; remove inherited publication placeholders.
    for sec in d.sections:
        for part in [sec.header,sec.footer]:
            for p in part.paragraphs:p.clear()
        sec.header_distance=Inches(.45);sec.footer_distance=Inches(.45)
        foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        if name not in d.styles:
            d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style=d.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11 if name=='Normal' else 17 if name=='Title' else 13 if name=='Heading 1' else 11)
        style.paragraph_format.space_after=Pt(6)
        if name.startswith('Heading'):
            style.paragraph_format.space_before=Pt(12)
            style.paragraph_format.keep_with_next=True
    normal=d.styles['Normal'];normal.paragraph_format.line_spacing=1.04
    normal.paragraph_format.widow_control=True
    settings=d.settings.element
    for el in list(settings):
        if el.tag==qn('w:autoHyphenation'):settings.remove(el)
    hyph=OxmlElement('w:autoHyphenation');hyph.set(qn('w:val'),'false');settings.append(hyph)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');normal.element.get_or_add_rPr().append(lang)
    textlog=[]

    def p(t,style=None):
        paragraph=d.add_paragraph(t,style or 'Normal');textlog.append(t)
        if not style:paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return paragraph
    def h(t,level=1):return p(t,f'Heading {level}')
    def figure(name,caption,width=5.65):
        paragraph=d.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next=True
        paragraph.add_run().add_picture(str(REV/'figures'/f'{name}.png'),width=Inches(width))
        pic=paragraph._p.xpath('.//wp:docPr')[0];pic.set('descr',caption)
        cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_together=True
        cap.paragraph_format.space_after=Pt(9)
        for r in cap.runs:r.font.size=Pt(9)
    def table(caption,headers,rows,widths):
        cap=p(caption,'Caption');cap.paragraph_format.keep_with_next=True;cap.paragraph_format.keep_together=True
        tab=d.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
        for cell,width in zip(tab.rows[0].cells,widths):cell.width=Inches(width)
        for cell,t in zip(tab.rows[0].cells,headers):cell.text=str(t)
        for row in rows:
            cells=tab.add_row().cells
            for cell,t,width in zip(cells,row,widths):cell.text=str(t);cell.width=Inches(width)
        for row_i,row in enumerate(tab.rows):
            pr=row._tr.get_or_add_trPr()
            cant=OxmlElement('w:cantSplit');pr.append(cant)
            if row_i==0:
                repeat=OxmlElement('w:tblHeader');pr.append(repeat)
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    v=OxmlElement(f'w:{side}');v.set(qn('w:w'),'75');v.set(qn('w:type'),'dxa');margins.append(v)
                tcpr.append(margins)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    v=OxmlElement(f'w:{side}');v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
                tcpr.append(borders)
                if row_i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECEE');tcpr.append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                    para.paragraph_format.line_spacing=1.0
                    if caption.startswith('Table 1.'):
                        para.paragraph_format.keep_with_next=row_i<len(tab.rows)-1
                    for run in para.runs:run.font.size=Pt(9);run.font.bold=row_i==0
        d.add_paragraph().paragraph_format.space_after=Pt(0)

    def equation():
        paragraph=d.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math=OxmlElement('m:oMath')
        # Editable native Word math, with structured superscript for L2.
        for s in ['L = MSE + 0.5 mean[ReLU(∂f/∂T) + ReLU(−∂f/∂P)] + 10⁻⁴ ∑ ']:
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);math.append(r)
        sup=OxmlElement('m:sSup');base=OxmlElement('m:e');power=OxmlElement('m:sup')
        for parent,s in [(base,'θ'),(power,'2')]:
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);parent.append(r)
        sup.append(base);sup.append(power);math.append(sup);paragraph._p.append(math)

    title='CeresPINN Monotonicity Regularization and Validation in a Hybrid Maize Climate Simulation Prototype'
    p(title,'Title')
    p('David Adrian Lucano Nieves¹ and Geraldine Daniela Rojas Villegas¹')
    p('¹ Universidad Nacional de Trujillo, Escuela de Ingeniería de Sistemas, Av. Juan Pablo II s/n, Trujillo 13011, Peru')
    p('dalucanoni@unitru.edu.pe; gdrojasvi@unitru.edu.pe')
    h('Abstract',2)
    p(f'Hybrid crop applications require evaluation that distinguishes learned prediction from programmed scenario responses. We audit CeresPINN, a monotonicity-regularized neural predictor combined with a deterministic crop simulator, and test a separate state-coupled process extension. The legacy dataset contains 36 independent annual maize-yield targets; scenario copies do not increase sample size. The archived checkpoint achieved RMSE 712.1 kg/ha on eight grouped held-out years, while linear baselines had lower errors. Expanding-window evaluation covers 20 distinct test years in five blocks. The regularized network obtained RMSE {scores["monotonic_all"]["rmse_kg_ha"]:.1f} kg/ha, compared with {scores["ridge_all"]["rmse_kg_ha"]:.1f} for Ridge; four Holm-adjusted comparisons gave p = 1. A separate Iowa county audit assesses baseline transfer rather than validation of the national checkpoint. Collinearity prevents causal interpretation of legacy scenario sensitivities. Public USDA Bushland observations support calibration of six crop parameters using four seasons and evaluation using 2016 and 2018. Across eight evaluation field-seasons, coupled-process dry-grain RMSE was 3248.1 kg/ha versus 1604.2 for weather–irrigation Ridge. Daily biomass and evapotranspiration were partly reproduced, but grain transportability and soil-water skill were weak despite numerical water closure. The contribution is an auditable evaluation framework with an externally sourced failure analysis. Neither the original management multipliers nor the new process extension supports agronomic prescriptions or an operational digital-twin claim.')
    p('Keywords: maize yield; monotonicity regularization; temporal validation; hybrid simulation; reproducibility.').alignment=WD_ALIGN_PARAGRAPH.LEFT

    h('1 Introduction')
    p('Drought and extreme heat can reduce maize production through both yield losses and reductions in harvested area [1]. Global climate–crop projections show that the magnitude and timing of impacts depend on emissions, location, and crop-model assumptions [2]. Regional work in Northeast China likewise reports substantial risks and evaluates adaptation with explicit crop-model forcing [3]. These findings motivate climate-aware crop applications, but they do not validate the predictions of a newly developed software system.')
    p('Process-based crop models and machine learning can complement each other. Shahhosseini et al. incorporated crop-model outputs into yield prediction in the US Corn Belt [4], while Cheng et al. examined county-level early prediction with multiple indicators [5]. Digital-twin reviews emphasize data integration, model interoperability, and observation-driven updating [6], [7], [8]. A visualization or a scenario interface alone does not establish an operational digital twin. CeresPINN currently has no continuous field-observation assimilation or validated bidirectional connection to a physical crop.')
    p('Standard physics-informed neural networks constrain learning with governing equations [9]. Hybrid agricultural approaches have also combined process models and deep learning [10]; AgriPINN embeds a differentiable crop-growth constraint and was initially disseminated as a preprint [11]. PINN-SM explicitly models soil-moisture dynamics [12]. CeresPINN instead penalizes local signs of yield derivatives with respect to temperature and precipitation. Its separate daily simulator does not provide a Richards-equation residual, a mass-conservation loss, or simulated-state inputs to the yield network. We therefore retain CeresPINN as the software name and describe the trained component as a monotonicity-regularized network.')
    p('The objective is to determine which conclusions the implemented system and its available data support. RQ1 asks whether monotonicity regularization improves prediction relative to an otherwise comparable MLP and simple baselines under grouped and future-year evaluation. RQ2 asks whether results persist when time trend and redundant scenario proxies are separated. RQ3 asks what an internally held-out county experiment can establish about regional baseline transfer. RQ4 asks how programmed management rules and soil profiles affect the simulator and whether those states reach the yield predictor. These questions replace claims of operational twin performance, causal sensitivity, and optimized agronomic management.')
    p('The contribution is a checkpoint-grounded audit with distinct temporal and geographic estimands, a reduced-input and detrended comparison, and a reproducible correction of data-semantic and inference-provenance errors. No new regularization algorithm is claimed. The study also identifies concrete requirements for a future coupled model: measured intermediate crop states, locally calibrated soil properties, identifiable climate inputs, and independent prospective validation. Digital-twin wheat experiments [13] and soil-map evaluation [14] illustrate the type of observational evidence that this prototype still lacks; the broader sustainability implications require equally explicit limits [15].')
    p('A separate extension uses public measured crop and water states to test an explicit state-to-yield pathway. This extension addresses whether numerical coupling and calibration alone are sufficient for temporal transportability. It does not replace the national neural checkpoint or establish a new PINN algorithm.')

    h('2 Materials and Methods')
    h('2.1 Datasets and independent units',2)
    p('The legacy yield lineage reports 22,160 USDA-NASS observations from 1990–2025, reduced to 36 annual national medians. The tracked processed long table contains 66,480 rows, comprising three scenario copies of the reported yield records. Reconstructing the training matrix by scenario-year medians yields 108 rows, but there are only 36 independent annual targets. All copies of a year remain in the same partition and metrics are computed after averaging scenario predictions within year. The conversion used throughout is 1 bushel/acre = 62.77 kg/ha. The current raw NASS download is header-only, so the original raw-record count is treated as recorded lineage rather than independently reconstructed raw evidence.')
    p('CMIP6 supplies a common experimental framework [16], and NASA NEX-GDDP-CMIP6 provides downscaled daily projections [17]. The legacy input assembly combines regional summaries, interpolation, and scenario templates. Its recorded lineage contains 98 regional climate rows; the exact original model–member–year extraction is not recoverable from the preserved raw files. The defaults for SSP5-8.5 include a 2.7 °C temperature anomaly, a −24% precipitation anomaly, and 520 ppm CO₂. These are internal scenario settings, not a probabilistic 2050 ensemble projection. SSP2-4.5 was absent from legacy training, and the proposed seven-GCM ensemble was not run for this study.')
    p('The regional extension uses a separate NASS–NOAA panel of 3,460 Iowa county-years, 99 counties, and coverage within 1990–2025. Duplicate county-years, missing values, invalid FIPS identifiers, and aggregate identifiers are rejected. A new public NASS Quick Stats [27] export contains 3,476 original records; 16 combined-county rows without individual county codes are excluded, and all 3,460 remaining yields match the frozen panel within rounding. We also preserve 297 complete NOAA Climate at a Glance [26] JSON responses: three variables per county, each with 36 April–September aggregates, source units, query URL, acquisition time, and SHA-256 checksum. This establishes current-source provenance for the regional reconstruction, not the original national raw acquisition or a new independent prediction cohort.')
    p('The new NOAA values agree with the frozen table through 2024. In 2025, 20 mean-temperature, 22 maximum-temperature, and 14 precipitation cells differ, by at most 0.167 °C, 0.167 °C, and 32.512 mm. Missing original raw responses prevent attributing the differences conclusively to an agency revision or earlier ingestion. Both versions are retained. The identical fixed held-county protocol is rerun on the reconstructed table as a source-version robustness check, with no change to test years, model settings, or the archived checkpoint.')
    p('A semantic audit found that NOAA CAG cdd means cooling degree days, whereas an earlier pipeline treated it as consecutive dry days. We exclude this variable from the county benchmark and correct its future export name to cooling_degree_days_c. Approximate GDD, heat-day counts derived from seasonal mean temperature, and temperature-only VPD are also excluded because they are not observed daily agronomic quantities. The NOAA request uses a six-month window ending in September, which represents April–September. The retained historical column names do not imply a May–September window.')
    figure('figure1_lineage','Figure 1. Data lineage and validation units. Scenario expansion preserves the same annual yield target and does not create independent observations.')

    h('2.2 Neural predictor and deterministic simulator',2)
    p('The archived neural checkpoint has seven inputs: year, temperature anomaly, precipitation anomaly, CO₂ concentration, heatwave_risk, seasonal precipitation, and seasonal_cdd. The last is constant in training. Four hidden layers contain 128 units each with Tanh activation and dropout. The yield head is supervised; the auxiliary sigmoid head has no observed target and does not feed back into yield. Its range penalty is identically zero by construction. The checkpoint contains 63,042 trainable parameters, a large capacity relative to 36 annual targets.')
    p('The loss combines normalized yield MSE, a local monotonicity penalty, and explicit L2 regularization. Derivatives are taken with respect to standardized temperature and seasonal-precipitation inputs. Positive scaling preserves derivative sign. The penalty discourages increasing yield with temperature anomaly and decreasing yield with precipitation within the sampled training domain; it is a modeling assumption, not a universal agronomic law or an extrapolation guarantee. Waterlogging, temperature optima, and interactions may violate these signs outside the intended regime.')
    equation()
    p('Here f is normalized predicted yield, T and P are the corresponding standardized temperature and seasonal-precipitation inputs, and θ denotes trainable weights. The explicit L2 coefficient is 0.0001. Adam additionally applies weight decay 0.00001; these are separate terms. Scaling parameters and target minima/maxima are estimated from training observations only in every new evaluation fold.')
    table('Table 1. Neural architecture and fixed training configuration.', ['Item','Setting'],[
        ['Input schema','7 legacy features; reduced experiments use temperature anomaly and seasonal precipitation'],
        ['Hidden layers','4 × 128 units; Tanh; dropout 0.05'],
        ['Yield head','128 → 64 → 1; ReLU between linear layers'],
        ['Auxiliary head','128 → 32 → 1; ReLU then sigmoid; no observed target'],
        ['Trainable parameters','63,042 in the legacy architecture'],
        ['Optimizer','Adam; learning rate 0.001; no scheduler'],
        ['Regularization','Monotonicity weight 0.5; explicit L2 0.0001; Adam weight decay 0.00001'],
        ['Training','300 epochs; batch size 64; CPU'],
        ['Seeds','Legacy repeatability: 0–8 and 42; rolling audit: 0, 1, 42'],
    ],[1.45,4.65])
    p('The deterministic simulator tracks thermal time and phenological stage, canopy development, root depth, water storage in three soil layers, evapotranspiration, and stress indices. It uses prescribed hydraulic profiles rather than locally measured or spatially mapped soil inputs. It does not solve the Richards equation. Soil moisture and phenology change the displayed daily trajectories but are not features of the archived yield network. RQ4 therefore tests a software pathway and an architectural disconnection, not physiological coupling.')

    h('2.3 Climate scenarios and scope of sensitivity',2)
    p('Downscaling does not remove all climate-model error, and precipitation extremes remain difficult to represent [18]. Bias-corrected regional data products such as the Senegal CMIP6 dataset [19] document a different and more traceable processing chain. Multivariate adjustment can matter when crop models depend on correlated climate variables [20]. CeresPINN did not perform an additional CDF-t, QDM, or MBCn correction, nor did it evaluate a full ScenarioMIP ensemble [21]. Such methods are contextual alternatives rather than steps implemented here.')
    p('The legacy Sobol calculation independently varies temperature anomaly, precipitation anomaly, and CO₂ over its specified numerical ranges while holding other inputs fixed. It describes the mathematical response of that particular checkpoint under a chosen sampling distribution. Because scenario proxies are dependent in the training data, these independent perturbations may leave the observed data manifold. We retain the indices only as a diagnostic of the legacy implementation and withdraw causal or physiological rankings. A constant input cannot receive an identifiable sensitivity index.')
    p('We compute Pearson correlations and variance inflation factors using the initial 1990–2005 training window, before the new future-year blocks. A feature with zero variance is excluded; exact linear dependence is reported as infinite VIF rather than a finite number. Reduced neural models remove CO₂, heatwave_risk, precipitation anomaly, and the constant dry-day input together, retaining temperature anomaly and seasonal precipitation. This avoids claiming that deleting only CO₂ resolves the scenario confounding.')

    h('2.4 Evaluation designs and statistical analysis',2)
    p('The archived checkpoint is evaluated on its stored grouped-random-year split: 28 training years and eight test years (1991, 1992, 1998, 2001, 2003, 2004, 2011, and 2024). This is historical interpolation with complete-year separation, not a prospective forecast. Baseline coefficients and scalers are fitted using the same training years. Legacy annual bootstrap intervals use 20,000 resamples and assume exchangeable annual errors; they do not correct temporal dependence. Non-significant residual-mean or distribution tests are not equivalence tests.')
    p('A separate chronological holdout trains on 1990–2017 and tests 2018–2025. Ten neural seeds are averaged after training-only normalization. In the new rolling-origin audit, the initial training period is 1990–2005; successive test blocks are 2006–2009, 2010–2013, 2014–2017, 2018–2021, and 2022–2025. Each model is refitted using only years preceding its block. Thus every one of the 20 test years has a prediction made without access to that year or later years during its fit. The blocks are evaluated once with fixed hyperparameters, not used for model selection.')
    p('The rolling audit compares an OLS year trend, linear regression with all legacy features, Ridge with all features (alpha = 1), Ridge with the two climate features, an unregularized MLP, a regularized network with all features, a regularized climate-only network, and a detrended climate-only regularized network. Neural predictions average seeds 0, 1, and 42. For the detrended model, a training-only OLS year trend is removed from training yield, the network learns residuals from the two climate inputs, and the extrapolated trend is added back at prediction time. This separates secular trend from the neural climate response; it does not eliminate uncertainty about the trend\u2019s cause.')
    p('RMSE, MAE, R², and prediction-minus-observation bias are computed on annual predictions. Descriptive 95% intervals use a circular moving-block bootstrap with blocks of three consecutive years and 20,000 draws. Block length is fixed and the small sample still limits coverage. Four comparisons contrast the all-input regularized model with the trend, linear regression, Ridge, and unregularized MLP. Absolute-error differences are averaged within each four-year test block. An exact two-sided sign randomization over five blocks produces nominal p-values, followed by Holm adjustment across the four comparisons. Training windows overlap and five error blocks provide little power; the p-values are diagnostic, not proof of equivalence or independent replication.')
    p('For the geographic extension, county FIPS codes are shuffled with seed 42 and partitioned into five disjoint groups. Each fit uses other counties through 2017 and predicts held counties during 2018–2025. No county contributes both training and test observations within a fit. The combined test set has 696 county-years but only eight distinct climate years; county-year rows are not treated as independent annual replicates. We compare fixed Ridge models using year only, climate only, and year plus climate. No baseline is chosen or refitted for deployment using these test results. Intervals resample complete year clusters in three-year blocks and recompute pooled errors. Random county groups do not establish independence from spatially correlated neighboring counties.')
    p('The additional national and county analyses are retrospective internal audits. Although each model fit respects its temporal or geographic exclusion, the data collection and earlier results were already known to the developers. Neither the rolling analysis nor the Iowa extension constitutes a preregistered, untouched external evaluation.')
    p('A further geographic stress test excludes each of the nine official Iowa agricultural statistics districts [31] in turn. Membership comes directly from the preserved NASS CSV and is constant across the 99 counties and source years. Every fit uses other districts through 2017 and predicts the excluded district during 2018–2025. The same three fixed Ridge specifications and training-only scaling are used. The district protocol was saved before fitting; no boundaries, hyperparameters, or spatial buffer were chosen from its scores. These geographic blocks are more restrictive than randomly shuffled counties, but share only eight weather years and remain within one state. A zero-width district boundary is not an autocorrelation buffer.')

    h('2.5 Management demonstrations and reproducibility',2)
    p('The Iowa–SSP5-8.5–2050 demonstration uses explicit internal forcing settings, a medium-cycle cultivar, 180 kg N/ha, and an initial soil-moisture setting of 60%. The baseline irrigation multiplier is 0.96; the nominal 75% strategy uses 1.03 and the short-cycle cultivar uses 0.88. Their ratios mechanically determine yield changes after the neural output. Moving planting 14 days earlier affects the daily simulator but not the yield-network inputs. These comparisons are rule verification, not estimated treatment effects. No experimental calibration, uncertainty distribution, or optimization objective exists for the management multipliers, so no agronomic recommendation or calibrated confidence interval is reported.')
    p('The analysis checks that the trained checkpoint supplies a numerical prediction before generating scientific management artifacts. If inference falls back to a surrogate, the audit stops instead of recording fallback values as checkpoint evidence. The two 110-day spreadsheet demonstrations retained from the application represent a separate local Ames Norte–SSP3-7.0 interface export without a live backend; they are not the Iowa–SSP5-8.5–2050 neural experiment. Their values are presented solely to document deterministic trajectories.')

    h('2.6 External-source state-coupling experiment',2)
    p('The USDA-ARS Bushland collection [28], [30] supplies six maize seasons (1989, 1990, 1994, 2013, 2016, and 2018) at one Texas station (35.186714° N, 102.094189° W; 1,170 m). Original CC0 workbooks and API metadata are preserved with official MD5 and independently calculated SHA-256 checksums. Extraction retains source files, worksheets, and row numbers. Extracted tables contain 104 combine-harvest subsamples, 759 growth records, 7,094 daily water records, and 27,994 soil-water observations. Only calendar-matched field-seasons enter simulations: 20 in total, 19 with combine yield. Plot and daily records do not create independent seasons. Dry grain in kg/ha is analyzed separately from standard-moisture yield; the 2016 east-field kg/m² values are explicitly converted. Source date inconsistencies are flagged rather than silently treated as phenological dates.')
    p('The extension is a conservative three-layer bucket (0–30, 30–60, and 60–150 cm), with thickness-weighted measured Pullman soil properties. Daily grass-reference evapotranspiration follows FAO-56 [29] using measured temperature, radiation, humidity, and 2-m wind. Unmeasured 2013 pressure is derived from elevation. Missing lysimeter precipitation uses the station gauge; missing irrigation uses the recorded event log, with absence of an event explicitly assumed zero; missing dew is assumed zero and flagged. Complete growing-season weather is required. A preplant soil profile initializes the planting-day state when available, with its age reported; otherwise field capacity is used. This approximation can omit intervening preplant water fluxes. Future soil or crop observations are never assimilated.')
    p('Thermal time determines LAI and rooting depth; layer water availability limits transpiration and intercepted-radiation biomass growth. Final dry grain equals simulated biomass times a fitted harvest index. Thus soil and applied water affect the yield pathway directly, without year or scenario multipliers. The algebraic balance is ΔS = precipitation + irrigation + dew − actual ET − deep drainage − runoff. This bucket is not a Richards solver, and exact accounting does not establish physical accuracy. Subsurface drip is applied to the second layer for east fields from 2013; sprinkler inputs enter the surface layer.')
    p('Six parameters (maximum LAI, thermal time to peak and senescence, maximum crop coefficient, radiation-use efficiency, and harvest index) are jointly fitted by bounded least squares using only 1989, 1990, 1994, and 2013. Residual scales are 2 LAI units, 10,000 kg/ha biomass, 3 mm/day ET, 0.05 m³/m³ soil water, and 3,000 kg/ha dry grain; each target family has equal total squared weight within a field-season. Fixed defaults and a standardized Ridge baseline (alpha = 10; rain, irrigation, reference ET, GDD, radiation, and planting DOY) are retained. Both models use observed forcing from planting through DOY 320; this is a conditional full-season simulation, not an early-season forecast. Evaluation uses all four fields in 2016 and 2018, with no parameter or algorithm changes after scoring. The frozen protocol preceded fitting, but some source endpoints were previewed during extraction, so this is retrospective rather than preregistered validation. NW–SW irrigation contrasts are confounded with field and cannot identify causal treatment effects.')

    h('3 Results')
    h('3.1 Archived checkpoint and chronological holdout',2)
    p('The archived checkpoint reproduces annual RMSE 712.13 kg/ha, MAE 604.19 kg/ha, and R² 0.8251 on eight grouped test years (Figure 2). The annual bootstrap RMSE interval is 465.9–915.5 kg/ha and the R² interval spans −0.1071 to 0.9355. Mean bias is +46.6 kg/ha. The paired residual-mean test gives t(7) = 0.1734, p = 0.8672, and dz = 0.0613. These small-sample results do not establish calibrated equivalence or forecasting validity.')
    figure('figure2_hindcast','Figure 2. Archived-checkpoint predictions averaged across scenarios for eight grouped held-out years. The dashed line marks identity; these years are not a future-only test.')
    table('Table 2. Same-split legacy annual performance from regenerated artifacts.', ['Model','RMSE kg/ha','MAE kg/ha','R²'],[
        [r['model'].replace('_',' '),f'{float(r["rmse_kg_ha"]):.2f}',f'{float(r["mae_kg_ha"]):.2f}',f'{float(r["r2"]):.4f}'] for r in baseline
    ],[2.5,1.25,1.25,1.1])
    p('Linear regression and Ridge have lower point errors than the checkpoint. The annual bootstrap differences include zero for these two comparisons. The neural ablation is relevant to the effect of the chosen penalty within a fixed architecture; it cannot establish superiority to all alternatives. Ten-seed dispersion in Figure 3 describes optimization variability on one split, not ten independent data replications. Retraining values are specific to the recorded numerical environment and should be assessed using the prediction-level artifacts.')
    figure('figure3_multiseed','Figure 3. Ten neural seeds on the legacy grouped split. Points represent fitted models; black markers show the mean and one seed standard deviation. The test years are shared across seeds.')
    tm=temporal['cerespinn_ensemble_metrics'];ci=temporal['cerespinn_ensemble_bootstrap_95_ci']['rmse_kg_ha']
    p(f'The chronological 2018–2025 holdout yields ensemble RMSE {tm["rmse_kg_ha"]:.2f} kg/ha, MAE {tm["mae_kg_ha"]:.2f} kg/ha, and R² {tm["r2"]:.4f}; its annual-bootstrap RMSE interval is {ci[0]:.1f}–{ci[1]:.1f} kg/ha. This is substantially weaker than the grouped random-year result. A negative R² means the squared error exceeds that of the test-target mean used in the R² denominator; that test mean is a scoring reference, not a deployable forecast. The result limits claims of future predictive skill.')

    h('3.2 Rolling temporal validation and trend separation',2)
    table('Table 3. Expanding-window results over 20 distinct future test years.', ['Fixed model','RMSE kg/ha','MAE kg/ha','R²'],[
        [r['model'].replace('_',' '),f'{r["rmse_kg_ha"]:.1f}',f'{r["mae_kg_ha"]:.1f}',f'{r["r2"]:.3f}'] for r in audit['national']['scores']
    ],[2.8,1.15,1.15,1.0])
    r=scores['monotonic_all'];ridge=scores['ridge_all'];dt=scores['detrended_monotonic_climate']
    p(f'The regularized all-input network obtained RMSE {r["rmse_kg_ha"]:.1f} kg/ha (block-bootstrap 95% interval {r["block_bootstrap_95_ci"]["rmse_kg_ha"][0]:.1f}–{r["block_bootstrap_95_ci"]["rmse_kg_ha"][1]:.1f}), whereas Ridge obtained {ridge["rmse_kg_ha"]:.1f} kg/ha. Removing year and all scenario proxies produced a climate-only neural RMSE of {scores["monotonic_climate_only"]["rmse_kg_ha"]:.1f} kg/ha. Explicit detrending reduced that error to {dt["rmse_kg_ha"]:.1f} kg/ha, but did not surpass Ridge. Thus removing year alone does not solve the prediction problem; the temporal component must be modeled and evaluated explicitly.')
    p('Relative to the unregularized MLP, the regularized network had lower pooled RMSE, but the four fold-level randomization tests all had Holm-adjusted p = 1.000. The mean paired MAE difference was approximately −138.3 kg/ha against the MLP and +165.1 kg/ha against Ridge. These are descriptive effect magnitudes, and positive values favor the comparator. There is insufficient evidence for a superiority claim after accounting for the small number of evaluation blocks and multiplicity. RQ1 is therefore answered with improved point performance against this neural ablation, alongside continued inferiority in point error to simpler linear baselines.')
    figure('figure5_rolling_origin','Figure 4. Predictions made from training windows preceding each four-year block. Vertical dotted lines separate evaluation blocks. The curves compare the regularized network with Ridge and a detrended climate network.')

    h('3.3 Regional transfer to held counties',2)
    table('Table 4. Geographic future-year tests using newly acquired agency sources. Both designs predict 696 county-years during 2018–2025; they are not independent cohorts.', ['Geographic holdout','Ridge inputs','RMSE kg/ha','MAE kg/ha','R²'],[
        [design,r['model'].replace('_',' '),f'{r["rmse_kg_ha"]:.1f}',f'{r["mae_kg_ha"]:.1f}',f'{r["r2"]:.3f}'] for design,results in [('Five county groups',reacquired),('Nine districts',districts)] for r in results['scores']
    ],[1.4,1.7,1.0,1.0,1.0])
    p('Across 696 held county-years, the newly reacquired table yields trend-only Ridge RMSE 1,173.6 kg/ha, compared with 1,268.9 for year plus climate and 2,991.8 for climate only. Relative to the frozen inputs, the two climate-model RMSE changes are only +0.243 and +0.222 kg/ha; trend-only predictions are unchanged. The source-version differences therefore do not reverse the baseline ordering. These results establish regional source fidelity and a geographically excluded internal benchmark, not validation of the national neural checkpoint at county scale. Shared weather, random rather than contiguous county groups, and only eight test years remain limitations of RQ3.')
    p('Excluding complete agricultural districts increases RMSE to 1,224.8 kg/ha for trend only and 1,302.9 for trend plus climate; climate-only RMSE is 2,996.0. Trend-only R² declines from 0.115 in random county groups to 0.036 across excluded districts. The corresponding year-block RMSE intervals are 1,023.2–1,389.3 and 1,066.4–1,521.5 kg/ha for the trend-only and trend-plus-climate district models. These descriptive intervals overlap and do not establish a significant model difference. Geographic exclusion is therefore more demanding, while the available seasonal climate inputs still fail to improve point error over the trend baseline. This result cannot be transferred to the national neural checkpoint or to Peruvian fields.')

    h('3.4 Collinearity and sensitivity diagnostics',2)
    p('On the initial training window, precipitation anomaly, CO₂, and heatwave_risk form a rank-deficient predictor set, so their individual VIFs are infinite. seasonal_cdd is constant and is excluded from VIF and interpretable sensitivity rankings. The reduced temperature-and-precipitation design removes those scenario proxies; this changes the model estimand and requires retraining, rather than reinterpreting the original checkpoint. The correlation structure is shown in Figure 5. RQ2 indicates that apparent CO₂ importance cannot be separated from scenario coding in the legacy design.')
    figure('figure6_confounding','Figure 5. Correlations in the initial 1990–2005 training window. Near-perfect scenario-proxy relations explain the rank-deficient full design. Correlation is not a causal effect.')
    p('The legacy independent-input Sobol indices are S1 = 0.0161 for temperature anomaly, 0.0044 for precipitation anomaly, and 0.9771 for CO₂. These values describe numerical variance under artificial independent perturbations and are not retained as a physiological ranking. The earlier mean-imputation test increases RMSE by 188.2% when year is replaced by its training mean, whereas 5% input noise has little effect. Mean imputation is a perturbation experiment, not retraining without year; Table 3 provides the latter comparison. Neither a high Sobol index nor low noise sensitivity supports causal attribution or reliable extrapolation to 2050.')

    h('3.5 Deterministic management and daily trajectories',2)
    table('Table 5. Recomputed software demonstration for Iowa SSP5-8.5 in 2050.', ['Programmed setting','Yield kg/ha','Change %','Irrigation mm'],[
        [r['scenario'].replace('_',' '),r['yield_kg_ha'],f'{float(r["delta_pct"]):.2f}',r['irrigation_mm']] for r in manage
    ],[2.65,1.15,1.1,1.2])
    p('The nominal 75% irrigation strategy changes yield by approximately +7.29% because 1.03/0.96 − 1 = 0.0729. A short cultivar applies a 0.88 factor; earlier planting leaves the neural yield unchanged. The factors have no fitted treatment-effect uncertainty, so reporting a confidence interval for their agronomic benefit would be unjustified. The four prescribed soil profiles can produce different daily water stress while retaining the same neural base yield. RQ4 therefore confirms the separation of pathways, not successful crop–soil feedback or an adaptation recommendation.')
    figure('figure4_daily_demonstration','Figure 6. Separate local interface demonstration for Ames Norte under SSP3-7.0 over 110 days: (a) reference evapotranspiration, precipitation, and irrigation; (b) biomass and LAI. These deterministic exports were produced without a live backend and are not the 2050 checkpoint experiment.',width=5.5)
    p('The accompanying S1 and S2 spreadsheets preserve the original daily interface exports. Their water-flux totals are 316.7 mm precipitation, 156.0 mm irrigation, and 729.6 mm reference evapotranspiration over 110 days; final reported biomass is 25,952 kg/ha. Reference evapotranspiration is atmospheric demand and cannot be substituted for actual crop evapotranspiration in a conservation equation. These outputs have no independent field validation and their data provenance is separate from the neural benchmark.')
    water=ext['water_balance_from_rounded_output']
    p(f'A distinct audit of rounded backend daily records finds a cumulative water-accounting residual of {water["sum_residual_mm"]:.3f} mm and mean absolute daily residual {water["mean_abs_daily_residual_mm"]:.3f} mm over {water["n_days"]} days. Because this check uses rounded exported states, it does not isolate numerical integration error from export rounding. It prevents a claim of demonstrated exact conservation; internal full-precision closure and observed soil-moisture validation remain necessary.')

    h('3.6 Measured states and external-source temporal evaluation',2)
    table('Table 6. Bushland dry-grain errors; calibration has 11 observed field-seasons and evaluation has eight, nested in two years.', ['Model','Calibration RMSE kg/ha','Evaluation RMSE kg/ha','Evaluation R²'],[
        [name.replace('_',' '),f'{v["calibration"]["rmse"]:.1f}',f'{v["evaluation"]["rmse"]:.1f}',f'{v["evaluation"]["r2"]:.3f}'] for name,v in field_metrics['models'].items()
    ],[2.4,1.3,1.3,1.1])
    p('The fitted parameters were maximum LAI 4.331, peak thermal time 1,262.3 °C·day, senescence 1,979.1 °C·day, crop coefficient 1.205, radiation-use efficiency 2.528 g/MJ PAR, and harvest index 0.399. Optimization converged with all six parameters inside the bounds; numerical Jacobian rank was six, which does not establish biological identifiability. Calibration improved grain error, but the fitted process transferred poorly to the two later seasons (Table 6). The fixed default process happened to perform better on these seasons; it was not selected for deployment after viewing the result. None of the four baselines had positive evaluation grain R².')
    figure('figure7_external_states','Figure 7. Bushland evaluation in 2016 and 2018: measured versus coupled-process predictions for (a) field-season dry grain and (b) field-day biomass. Daily points are repeated measurements within eight field-seasons, not independent trials.',width=5.5)
    p(f'Evaluation RMSE was {field_metrics["daily_states"]["LAI"]["evaluation"]["rmse"]:.3f} for LAI (49 field-days), {field_metrics["daily_states"]["biomass_kg_ha"]["evaluation"]["rmse"]:.1f} kg/ha for biomass (49), {field_metrics["daily_states"]["ET_mm"]["evaluation"]["rmse"]:.3f} mm/day for actual ET (1,455), and {field_metrics["daily_states"]["theta_m3_m3"]["evaluation"]["rmse"]:.4f} m³/m³ for layer soil water (333 layer-days). Biomass and ET R² were 0.749 and 0.759; soil-water R² was −0.191. These are descriptive repeated-measurement scores. The maximum absolute full-precision daily mass residual was {field_metrics["maximum_absolute_water_residual_mm"]:.2e} mm. This verifies numerical accounting in the new extension while leaving poor soil-state and grain transportability visible; it does not repair the separate legacy simulator or validate the national network.')

    h('4 Discussion')
    h('4.1 What the evaluation supports',2)
    p('The strongest supported conclusion is about evaluation design. A grouped random-year holdout can report high R² while future-only evaluation is weak. Enlarging the temporal assessment to 20 years improves coverage but does not create a large independent sample or a new external cohort. Regularization reduces point error against an equivalent unconstrained network in these experiments; the available evidence does not establish superiority to Ridge or a broadly generalizable agronomic advantage.')
    p('The work differs from process-model feature coupling [4] because its daily simulated states do not enter the archived yield network. It differs from AgriPINN [11] and PINN-SM [12] because no crop differential equation or Richards residual constrains the neural solution. The county extension is a separate baseline transfer experiment, not a retroactive spatial validation of that neural checkpoint. This distinction is essential when interpreting the software as a hybrid simulation prototype rather than an operational digital twin.')
    p('The detrended comparison makes secular yield change explicit, but it does not identify its mechanism. Technological improvement, cultivar changes, management, and climate may all contribute. Extrapolating an OLS trend beyond the available period is itself an assumption. The poor climate-only results and rank-deficient scenario proxies indicate that the current data cannot support robust attribution of long-term yield changes to CO₂, heat waves, or emissions scenarios.')

    h('4.2 Management interpretation and future coupling',2)
    p('Published adaptation studies show that planting periods and management can affect simulated crop outcomes [22], [23], and irrigation responses depend on water availability and regional conditions [24]. Their evidence cannot calibrate the CeresPINN multipliers by analogy. The present management module can be used to test software pathways, but it cannot quantify expected field benefits, optimize irrigation, or rank resilient cultivars. Its outputs establish programmed behavior rather than prescriptive agronomic effects.')
    p('The Bushland extension supplies a separate measured-state coupling experiment alongside the unchanged legacy system. Its negative transportability result shows why algebraic conservation and a fitted yield endpoint are insufficient. A common crop parameterization cannot resolve unmodeled cultivar, establishment, senescence, or management changes between these seasons. Field-capacity or stale preplant initialization and coarse layer averages also limit soil-water realism. Crop-model and machine-learning integration in data-limited environments [25] remains relevant, but predictive states must demonstrate skill outside calibration. Only two evaluation years and one site preclude reliable season-cluster uncertainty intervals or a claim of general agronomic validity; parameter uncertainty and multisite validation remain open.')

    h('4.3 Threats to validity and responsible use',2)
    table('Table 7. Remaining threats and the scope of the implemented corrections.', ['Threat','Correction applied','Remaining requirement'],[
        ['Small effective sample','20-year rolling audit; complete-year grouping; clustered intervals','Longer records and independently acquired external sites'],
        ['Scenario collinearity','Training-only correlation/VIF audit; reduced and detrended retraining','Identifiable observed climate drivers and consistent scenario forcing'],
        ['Geographic generalization','Held counties and nine official districts; raw agency sources reacquired','Spatial buffers, additional states, external neural validation'],
        ['Deterministic management','Rules and their ratios disclosed; no recommendation or calibrated effect claimed','Experimental calibration with treatment-effect uncertainty'],
        ['Uncoupled daily states','Legacy RQ4 corrected; separate measured-state extension tested','Better soil/grain transportability; no automatic legacy validation'],
        ['Reproducibility','Public Git snapshot, source files, hashes, scripts, seeds and environment','Complete original national raw lineage and archival DOI'],
    ],[1.25,2.45,2.4])
    p('The legacy national aggregation cannot represent cultivar, irrigation, county soil, harvested-area, pest, or disease heterogeneity. The Iowa extension now includes both random county exclusions and official district blocks with directly documented current-source reconstruction. It still covers one state, has no buffer between neighboring districts, and shares climate shocks across regions. Lost national raw records and unresolved original climate-template lineage remain separate limitations. Neither district blocking nor the single-station Bushland experiment establishes transfer to another agroecological population. There are no Peruvian field observations, and US results cannot be transferred directly to Peruvian maize systems.')
    p('The data do not support territorial vulnerability rankings, policy priorities, insurance pricing, credit allocation, or irrigation prescriptions. The prototype is suitable for research and software demonstration while these limitations remain. An operational digital twin would additionally need real observation updates, data-quality monitoring, an empirically assessed uncertainty model, and a validated feedback loop. Those functions are not established by the current interface or by the existence of a trained checkpoint.')

    h('5 Conclusions')
    p('CeresPINN integrates a monotonicity-regularized yield network with a deterministic crop and water-balance simulator in an auditable research prototype. The archived eight-year grouped result is reproducible, but it overstates future predictive ability if read as a prospective test. The expanded rolling analysis finds lower point error for Ridge than for the regularized network and no significant superiority in the four Holm-adjusted fold comparisons. Removing temporal and scenario proxies exposes weak climate-only prediction; explicit trend separation is more informative than using year implicitly.')
    p('The county and official-district exclusions add regional baseline evidence while preserving their distinction from national neural validation. Current agency sources reproduce all regional targets; the disclosed 2025 climate differences do not reverse baseline ordering. Collinearity diagnostics withdraw a causal interpretation of legacy Sobol rankings. The separately calibrated Bushland extension links water and crop states to grain and closes numerical water accounting, but later-season grain errors exceed Ridge and soil-water skill is weak. Measured-state coupling supplies a falsifiable pathway rather than proof of agronomic value. Original management multipliers remain uncalibrated. Multisite validation, original national raw-data recovery, and identifiable management effects remain necessary before climate-adaptation or operational claims.')

    h('Data and code availability',2)
    commit=audit['base_commit']
    publication_path=DEST/'PUBLICACION_GITHUB.json'
    if publication_path.exists():
        publication=json.loads(publication_path.read_text(encoding='utf-8'))
        availability=p(f'The public scientific revision is https://github.com/David42024/CeresPINN/commit/{publication["commit"]}. This Git commit contains code, frozen processed inputs, the archived neural checkpoint and metadata, environment versions, SHA-256 checksums, prediction-level audit tables, original public USDA/NOAA/Bushland source files and reproduction scripts. The manuscripts and private editorial response are excluded. Code is MIT licensed; source-data terms remain applicable. Original national NASS records and the original CMIP6 extraction remain unavailable. No archival DOI is claimed.')
        availability.alignment=WD_ALIGN_PARAGRAPH.LEFT
    else:
        p(f'The project repository is https://github.com/David42024/CeresPINN. The inspected base revision is {commit}; the present corrections are distributed as a separate local revision bundle and must not be attributed to that base commit. Code is licensed under MIT; underlying agency data remain governed by their source terms. The bundle contains frozen processed inputs, the archived neural checkpoint and metadata, environment versions, SHA-256 checksums, prediction-level audit tables, and the scripts used for these results. It does not contain a reconstructed original NASS download or a complete original CMIP6 extraction, and it is not represented as a publicly deposited DOI release.')
    p('From the bundle root, run python scripts/revision_scientific_audit.py for the new temporal and county experiments. Legacy audits use scripts/generate_q1_revision_artifacts.py, scripts/generate_q1_extended_audit.py, and scripts/run_temporal_validation.py. The historical-input integrity check is python scripts/reproduce.py --verify-only. Frozen versions for the regenerated analyses are in requirements-revision.txt and audit.json. Outputs are internal reproducibility artifacts, not independent validation data. S1 is the original daily water-flux spreadsheet and S2 the original biomass/LAI spreadsheet; both represent the separate interface demonstration described in Figure 6.')
    p('The Bushland experiment uses scripts/prepare_bushland_data.py followed by scripts/validate_bushland_coupling.py. The bundle includes the selected original public workbooks, acquisition metadata/checksums, processed observations and forcing, fixed protocol, fitted parameters, daily predictions, and field-season errors. Raw data access is also reproducible through scripts/acquire_bushland_data.py; the exact selected article/file identifiers are in the manifest. This external source is distinct from the incomplete legacy NASS/CMIP6 lineage. No evaluation-score-driven retuning was performed.')
    p('The regional source reconstruction includes the original public NASS CSV export, all 297 NOAA JSON responses and provenance records, the rebuilt county table, row-level differences, and source-version benchmark predictions. Run scripts/audit_reacquired_county.py to verify raw checksums and regenerate this separate comparison. Exact climate equality with the frozen table is explicitly false for the disclosed 2025 cells; numerical target agreement and source integrity are verified separately. This acquisition does not recover the lost national raw download or constitute a new untouched validation cohort.')
    p('Run scripts/audit_geographic_districts.py to reproduce the fixed nine-district exclusion experiment. Its frozen protocol, official grouping source, county-to-district assignments, split sizes, district-level errors, prediction-level records, and year-cluster intervals are included. No model was selected for deployment from the district scores.')
    h('Acknowledgments',2)
    p('This research received no specific grant from public, commercial, or nonprofit funding agencies. The authors developed the study using their own resources.')
    h('Generative AI disclosure',2)
    p('OpenAI Codex assisted with code inspection, reproducible analysis scripts, consistency checks, and English editing during preparation of this manuscript. The use of AI does not replace author review of code, source provenance, statistical interpretation, or the final text. The authors are responsible for reviewing and approving the submitted manuscript and its supporting artifacts.')
    h('Conflicts of interest',2)
    p('The authors declare no conflicts of interest related to this study.')

    # Check every in-text citation before adding the bibliography.
    used=set();citation_order=[]
    for t in textlog:
        for number in re.findall(r'\[(\d+)\]',t):
            number=int(number)
            if number not in used:citation_order.append(number)
            used.add(number)
    expected=set(range(1,len(REFERENCES)+1))
    if used!=expected:raise ValueError(f'Citation mismatch: missing={expected-used}; unknown={used-expected}')
    renumber={old:i+1 for i,old in enumerate(citation_order)}
    def replace_citations(t):
        return re.sub(r'\[(\d+)\]',lambda m:f'[{renumber[int(m.group(1))]}]',t)
    for paragraph in d.paragraphs:
        if re.search(r'\[\d+\]',paragraph.text):paragraph.text=replace_citations(paragraph.text)
    textlog[:]=[replace_citations(t) for t in textlog]
    h('References')
    for i,old in enumerate(citation_order,1):
        ref,link=REFERENCES[old-1]
        url=link if link.startswith('https://') else 'https://doi.org/'+link
        para=p(f'[{i}] {ref} {url}')
        para.paragraph_format.left_indent=Inches(.24);para.paragraph_format.first_line_indent=Inches(-.24)
        para.paragraph_format.space_after=Pt(6)
        para.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:run.font.size=Pt(10)
    d.core_properties.title=title
    d.core_properties.author='David Adrian Lucano Nieves; Geraldine Daniela Rojas Villegas'
    d.core_properties.subject='Scientific revision with temporal and county validation audits'
    out=DEST/'CeresPINN_articulo_corregido.docx';d.save(out)
    (DEST/'CeresPINN_articulo_corregido.txt').write_text('\n\n'.join(textlog),encoding='utf-8')
    for source,name in [('cerespinn_reporte_balance_hidrico.xlsx','S1_balance_hidrico.xlsx'),
                        ('cerespinn_reporte_biomasa_lai.xlsx','S2_biomasa_lai.xlsx')]:
        shutil.copy2(ROOT/'output'/source,DEST/name)
    (REV/'manuscript_integrity.json').write_text(json.dumps({
        'original_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        'revised_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),
        'citation_numbers':sorted(used),'bibliography_count':len(REFERENCES),
        'reference_renumbering':renumber,
        'tables':len(d.tables),'figures':len(d.inline_shapes),
        'original_preserved':True},indent=2),encoding='utf-8')
    print(out)


if __name__=='__main__':main()
