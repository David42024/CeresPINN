"""Convierte el artículo suministrado a Markdown y Word sin revisar sus datos."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('C:/Users/USERJSSV/.codex/attachments/cb49d59f-2206-4227-914e-eee33ddfc76c/Texto pegado.txt')

def convertir_articulo(origen=SOURCE, destino=ROOT / 'docs'):
    destino.mkdir(exist_ok=True)
    lines = origen.read_text(encoding='utf-8-sig').splitlines()
    blocks = [('title', lines[0])]
    i = 1
    special = {'Resumen', 'Declaraciones', 'Referencias', 'Material Suplementario', 'Checklist de Cumplimiento Q1'}
    while i < len(lines):
        s = lines[i]
        if not s.strip():
            i += 1
            continue
        if s == 'Hojas':
            i += 1
            continue
        if s == r'\begin{equation}':
            arr = [s]
            i += 1
            while i < len(lines):
                arr.append(lines[i])
                i += 1
                if arr[-1] == r'\end{equation}':
                    break
            blocks.append(('math', '\n'.join(arr)))
            continue
        if s in {'plain', 'bash'}:
            lang = 'text' if s == 'plain' else 'bash'
            arr = []
            i += 1
            while i < len(lines) and not lines[i].startswith('Tabla '):
                arr.append(lines[i])
                i += 1
            blocks.append(('code', (lang, '\n'.join(arr).rstrip())))
            continue
        if '\t' in s:
            rows = []
            while i < len(lines) and '\t' in lines[i]:
                row = lines[i]
                i += 1
                # Reúne símbolos y exponentes que el pegado separó en varias líneas.
                while i < len(lines) and (lines[i].startswith('  ') or re.fullmatch(r'[−⁻]?\d+|data|phys|reg|\u200b', lines[i].strip())):
                    row += ' ' + lines[i].strip()
                    i += 1
                    if '\t' in lines[i - 1]:
                        break
                rows.append([c.strip() for c in row.split('\t')])
            blocks.append(('table', rows))
            continue
        match = re.match(r'^(\d+(?:\.\d+)*)(?:\.)?\s+\S', s)
        if match and not re.match(r'^\d+:', s):
            blocks.append(('heading', (min(3, match[1].count('.') + 1), s)))
        elif s in special:
            blocks.append(('heading', (1, s)))
        elif s.startswith(('Figura ', 'Tabla ', 'Algoritmo ', 'Código S')):
            blocks.append(('caption', s))
        else:
            blocks.append(('paragraph', s))
        i += 1
    md = []
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(0.75)
    for style in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[style].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(20)
    code_style = doc.styles.add_style('Texto técnico', 1)
    code_style.font.name = 'Consolas'
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_after = Pt(2)
    for kind, data in blocks:
        if kind == 'title':
            md.append('# ' + data)
            doc.add_paragraph(data, 'Title')
        elif kind == 'heading':
            level, text = data
            md.append('#' * (level + 1) + ' ' + text)
            doc.add_heading(text, level)
        elif kind == 'caption':
            md.append('**' + data + '**')
            p = doc.add_paragraph(data, 'Caption')
            p.paragraph_format.keep_with_next = True
            p.runs[0].bold = True
        elif kind in {'code', 'math'}:
            lang, text = data if kind == 'code' else ('latex', data)
            md.append('```' + lang + '\n' + text + '\n```')
            for line in text.splitlines():
                doc.add_paragraph(line, 'Texto técnico')
        elif kind == 'table':
            n = max(map(len, data))
            rows = [r + [''] * (n - len(r)) for r in data]
            md.append('\n'.join(['| ' + ' | '.join(c.replace('|', '\\|') for c in rows[0]) + ' |',
                                  '| ' + ' | '.join(['---'] * n) + ' |'] +
                                 ['| ' + ' | '.join(c.replace('|', '\\|') for c in r) + ' |' for r in rows[1:]]))
            table = doc.add_table(rows=0, cols=n)
            table.style = 'Table Grid'
            for idx, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(8)
                            run.bold = idx == 0
                    props = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E7E6E6' if idx == 0 else 'FFFFFF')
                    props.append(shade)
                if idx == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[-1]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        else:
            md.append(data)
            p = doc.add_paragraph(data)
            if len(data) > 200:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    md_path = destino / 'Articulo_CeresPINN.md'
    word_path = destino / 'Articulo_CeresPINN.docx'
    md_path.write_text('\n\n'.join(md) + '\n', encoding='utf-8')
    doc.save(word_path)
    # Verificación de integridad: todo el texto no estructural debe conservarse.
    original = ''.join(s for s in lines if s not in {'Hojas', 'plain', 'bash'})
    combined = ''.join(''.join(''.join(row) for row in data) if kind == 'table'
                       else (data[1] if kind in {'heading', 'code'} else data)
                       for kind, data in blocks)
    normalize = lambda s: re.sub(r'\s+', '', s)
    for s in lines:
        if s.strip() and s not in {'Hojas', 'plain', 'bash'}:
            assert normalize(s) in normalize(combined), repr(s)
    reread = Document(word_path)
    assert len(reread.tables) == sum(k == 'table' for k, _ in blocks)
    print(f'Contenido verificado. {len(blocks)} bloques; {len(reread.tables)} tablas.')
    print(md_path)
    print(word_path)
    return md_path, word_path

if __name__ == '__main__':
    convertir_articulo()
